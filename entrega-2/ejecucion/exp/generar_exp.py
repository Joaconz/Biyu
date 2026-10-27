"""Genera los entregables de la ejecución de Exportar CSV (CP-EXP-001 a 075) de la Entrega 2, con el mismo
formato que la Entrega 1 (entrega-1/fuentes/generar.py): Markdown + planillas .xlsx con la plantilla de la cátedra.

Fuentes: ../../casos/exp.md (casos), resultados-a/b/c.json (los resultados de casos-a/b/c.mjs),
obtenido.json (qué se observó en cada caso) y defectos.json (DEF-030 a DEF-032).

Uso (desde la raíz del repo, con Python 3 y openpyxl; en Windows, con PYTHONUTF8=1):
  python entrega-2/ejecucion/exp/generar_exp.py
Escribe en entrega-2/: 04-ejecucion-exp, 05-reportes-de-defectos-exp y 06-reporte-de-ejecucion-exp (.md y .xlsx).
"""
import json
import re
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
OUT = HERE.parent.parent  # entrega-2/
MD = (OUT / 'casos' / 'exp.md').read_text(encoding='utf-8')
RES = {}
for f in ('resultados-a.json', 'resultados-b.json', 'resultados-c.json'):
    for r in json.loads((HERE / f).read_text(encoding='utf-8')):
        RES[r['id']] = r  # una corrida posterior reemplaza a la anterior
GOT = json.loads((HERE / 'obtenido.json').read_text(encoding='utf-8'))
DEFECTOS = json.loads((HERE / 'defectos.json').read_text(encoding='utf-8'))

PROYECTO = 'Biyu'
FECHA = '2026-10-27'
COMMIT = '1c72d0f'
CREADO_POR = 'Equipo Biyu (casos de Exportar CSV, entrega-2/casos/exp.md)'
EJECUTADO_POR = 'Claude Code (IA), supervisado por Santiago'
ENTORNO = (f'Local: Vite (http://127.0.0.1:5173) + Supabase local · main {COMMIT} · Chromium (Playwright 1.63), '
           'es-AR, America/Argentina/Buenos_Aires, reloj del navegador fijado en 09/10/2026')

LABEL = PatternFill('solid', fgColor='FFFFFFCC')
EXEC = PatternFill('solid', fgColor='FFC6EFCE')
HEAD = PatternFill('solid', fgColor='FF1B4D3E')
STATUS_FILL = {'PASSED': 'FFC6EFCE', 'FAILED': 'FFFFC7CE', 'BLOCKED': 'FFFFEB9C', 'DUDA': 'FFDDEBF7', 'NO EJECUTADO': 'FFD9D9D9'}
THIN = Side(style='thin', color='FFBFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
BOLD = Font(bold=True)
WHITE = Font(bold=True, color='FFFFFFFF')


# ---------------------------------------------------------------- casos (parseo de casos/exp.md)
def parse_cases():
    cases = []
    for blk in re.split(r'^### ', MD, flags=re.M)[1:]:
        m = re.match(r'(CP-EXP-\d{3}) — (.+)', blk)
        if not m:
            continue
        cid, titulo = m.group(1), m.group(2).strip()

        def field(name):
            x = re.search(r'\| ' + re.escape(name) + r' \| (.+?) \|', blk)
            return x.group(1).strip() if x else ''

        pre = re.search(r'\*\*Pre-requisitos:\*\* (.+)', blk).group(1).strip()
        datos = re.search(r'\*\*Datos de prueba:\*\* (.+)', blk).group(1).strip()
        post = re.search(r'\*\*Post-condición:\*\* (.+)', blk)
        pasos = [(re.sub(r'\s+', ' ', a.strip()), re.sub(r'\s+', ' ', b.strip()))
                 for a, b in re.findall(r'^\| \d+ \| (.+?) \| (.+?) \|$', blk, flags=re.M)]
        pre_items = [p.strip() for p in re.split(r'(?=\bS\d+: )', pre) if p.strip()]
        datos_items = []
        for d in [p.strip() for p in re.split(r'(?=\bD\d+: )', datos) if p.strip()]:
            k, _, v = d.partition(': ')
            datos_items.append((k, v) if re.fullmatch(r'D\d+', k) else ('D', d))
        # El "Esperado" del caso: los resultados esperados de los pasos de la prueba, sin el ingreso a la app.
        skip = 0
        for i, (p, _) in enumerate(pasos):
            if p.startswith('Hacer click en') and 'Entrar' in p and 'barra' not in p:
                skip = i + 1
                break
        cases.append({
            'id': cid, 'titulo': titulo, 'prioridad': field('Prioridad'), 'tipo': field('Tipo'), 'canal': field('Canal'),
            'tecnica': field('Técnica'), 'trazas': field('Historia · criterio'), 'invariante': field('Invariante'),
            'par': field('Caso par'), 'automatizable': field('Automatizable'),
            'historias': ['US-47'], 'feliz': field('Tipo') == 'Positivo',
            'funcionalidad': 'Exportar movimientos a CSV (US-47) · ' + field('Historia · criterio').replace('US-47 · ', ''),
            'pre': pre_items, 'datos': datos_items, 'pasos': pasos, 'esperado_pasos': [(p, e) for p, e in pasos[skip:] if 'barra de direcciones' not in p and not p.startswith('Escribir `<APP>') and p != 'Presionar Enter.'],
            'post': post.group(1).strip() if post else '',
        })
    return cases


CASOS = parse_cases()
assert len(CASOS) == 75, len(CASOS)
BY_ID = {c['id']: c for c in CASOS}
CRITERIO = {c['id']: c['trazas'].replace('US-47 · ', '') for c in CASOS}


def status_of(cid):
    r = RES.get(cid)
    return r['status'] if r else 'NO EJECUTADO'


def defects_of(cid):
    return ', '.join(d['id'] for d in DEFECTOS if cid in d['casos'])


def esperado(c):
    return ' '.join(e for _, e in c['esperado_pasos'])


def obtenido(cid):
    r = RES[cid]
    parts = [GOT[cid]]
    if r.get('obs'):
        parts.append('— Observado en la corrida —')
        parts += r['obs']
    if r['status'] == 'FAILED' and r.get('note'):
        parts.append('Falla detectada: ' + r['note'])
    return '\n'.join(parts)


def text_height(text, chars_per_line):
    lines = sum(max(1, -(-len(part) // chars_per_line)) for part in str(text).split('\n'))
    return max(15, 15 * lines + 2)


def cell(ws, ref, value, fill=None, bold=False, merge=None):
    ws[ref] = value
    ws[ref].alignment = WRAP
    ws[ref].border = BOX
    if fill:
        ws[ref].fill = fill
    if bold:
        ws[ref].font = BOLD
    if merge:
        ws.merge_cells(merge)
    return ws[ref]


def tc_sheet(wb, c):
    """Una hoja por caso, con la misma grilla que 'TC - Template' (igual que en la Entrega 1)."""
    ws = wb.create_sheet(c['id'].replace('CP-', ''))
    for col, w in zip('ABCDEFG', (22, 46, 22, 46, 22, 40, 40)):
        ws.column_dimensions[col].width = w
    cell(ws, 'A1', 'Test Case ID', LABEL, True); cell(ws, 'B1', c['id'])
    cell(ws, 'C1', 'Descripción / Funcionalidad', LABEL, True); cell(ws, 'D1', c['funcionalidad'], merge='D1:E1')
    cell(ws, 'A2', 'Creado por', LABEL, True); cell(ws, 'B2', CREADO_POR)
    cell(ws, 'C2', 'Versión', LABEL, True); cell(ws, 'D2', '2', merge='D2:E2')
    cell(ws, 'A3', 'Reporte de ejecución 1', EXEC, True, merge='A3:B3')
    cell(ws, 'C3', 'Reporte de ejecución 2', EXEC, True); cell(ws, 'D3', 'Reporte de ejecución 3', EXEC, True, merge='D3:E3')
    st = status_of(c['id'])
    cell(ws, 'A4', f'{st} · {FECHA}', PatternFill('solid', fgColor=STATUS_FILL[st]), True, merge='A4:B4')
    cell(ws, 'A5', 'Prioridad (Alta, Media, Baja)', LABEL, True); cell(ws, 'B5', c['prioridad'])
    cell(ws, 'C5', 'Camino feliz', LABEL, True); cell(ws, 'D5', 'Sí' if c['feliz'] else 'No')
    cell(ws, 'A6', 'Técnica / Tipo', LABEL, True); cell(ws, 'B6', f"{c['tecnica']} · {c['tipo']}")
    cell(ws, 'C6', 'Trazabilidad', LABEL, True); cell(ws, 'D6', c['trazas'], merge='D6:E6')
    row = 8
    cell(ws, f'A{row}', 'S #', LABEL, True); cell(ws, f'B{row}', 'Pre-requisitos', LABEL, True)
    cell(ws, f'C{row}', 'D #', LABEL, True); cell(ws, f'D{row}', 'Datos de prueba', LABEL, True, merge=f'D{row}:E{row}')
    n = max(len(c['pre']), len(c['datos']), 1)
    for i in range(n):
        r = row + 1 + i
        if i < len(c['pre']):
            cell(ws, f'A{r}', i + 1); cell(ws, f'B{r}', c['pre'][i])
        if i < len(c['datos']):
            cell(ws, f'C{r}', i + 1); cell(ws, f'D{r}', f'{c["datos"][i][0]}:'); cell(ws, f'E{r}', c['datos'][i][1])
        pre = c['pre'][i] if i < len(c['pre']) else ''
        dato = c['datos'][i][1] if i < len(c['datos']) else ''
        ws.row_dimensions[r].height = max(text_height(pre, 42), text_height(dato, 40))
    row += n + 2
    cell(ws, f'A{row}', 'Título del CP', LABEL, True); cell(ws, f'B{row}', c['titulo'], merge=f'B{row}:E{row}')
    row += 2
    cell(ws, f'A{row}', 'Historias de Usuario', LABEL, True); cell(ws, f'B{row}', ', '.join(c['historias']))
    cell(ws, f'C{row}', 'Estado Automatización', LABEL, True)
    auto = 'Automatizable (candidato)' if c['automatizable'].startswith('S') else 'No automatizable'
    cell(ws, f'D{row}', f'{auto}. Ejecución de esta entrega asistida por un script de Playwright (Claude Code).', merge=f'D{row}:E{row}')
    row += 2
    cell(ws, f'A{row}', 'Pasos #', LABEL, True); cell(ws, f'B{row}', 'Detalle de pasos', LABEL, True, merge=f'B{row}:C{row}')
    cell(ws, f'D{row}', 'Resultado Esperado', LABEL, True, merge=f'D{row}:E{row}')
    for i, (paso, esp) in enumerate(c['pasos'], 1):
        r = row + i
        cell(ws, f'A{r}', i); cell(ws, f'B{r}', paso, merge=f'B{r}:C{r}'); cell(ws, f'D{r}', esp, merge=f'D{r}:E{r}')
        ws.row_dimensions[r].height = max(text_height(paso, 62), text_height(esp, 62))
    row += len(c['pasos']) + 2
    r = RES[c['id']]
    texto = obtenido(c['id'])
    cell(ws, f'A{row}', 'Resultado obtenido', EXEC, True); cell(ws, f'B{row}', texto, merge=f'B{row}:E{row}')
    ws.row_dimensions[row].height = max(30, text_height(texto, 120))
    cell(ws, f'A{row + 1}', 'Evidencia', EXEC, True); cell(ws, f'B{row + 1}', '\n'.join(r.get('evidence', [])) or '—', merge=f'B{row + 1}:E{row + 1}')
    cell(ws, f'A{row + 2}', 'Defectos', EXEC, True); cell(ws, f'B{row + 2}', defects_of(c['id']) or '—', merge=f'B{row + 2}:E{row + 2}')
    cell(ws, f'A{row + 3}', 'Notas del ejecutor', EXEC, True); cell(ws, f'B{row + 3}', (r.get('note') if r['status'] != 'FAILED' else '') or '—', merge=f'B{row + 3}:E{row + 3}')
    cell(ws, f'A{row + 4}', 'Entorno / ejecutó', EXEC, True); cell(ws, f'B{row + 4}', f'{ENTORNO}\n{EJECUTADO_POR}', merge=f'B{row + 4}:E{row + 4}')
    return ws


def index_sheet(wb, title):
    ws = wb.active
    ws.title = 'Índice'
    heads = ['Test case ID', 'Prioridad', 'Título', 'User story', 'Camino feliz', 'Tipo', 'Técnica', 'Status', 'Ejecutó', 'Defectos']
    ws.append([title]); ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
    ws.append(heads)
    for i in range(1, len(heads) + 1):
        c = ws.cell(row=2, column=i); c.fill = HEAD; c.font = WHITE; c.border = BOX; c.alignment = WRAP
    for c in CASOS:
        ws.append([c['id'], c['prioridad'], c['titulo'], c['trazas'], 'Sí' if c['feliz'] else 'No', c['tipo'], c['tecnica'], status_of(c['id']), 'Claude Code (IA)', defects_of(c['id'])])
        for i in range(1, len(heads) + 1):
            x = ws.cell(row=ws.max_row, column=i); x.border = BOX; x.alignment = WRAP
        ws.cell(row=ws.max_row, column=8).fill = PatternFill('solid', fgColor=STATUS_FILL[status_of(c['id'])])
    for col, w in zip('ABCDEFGHIJ', (14, 10, 58, 22, 12, 11, 24, 12, 18, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'


# ---------------------------------------------------------------- métricas
def metrics():
    st = Counter(status_of(c['id']) for c in CASOS)
    by = lambda key, vals: OrderedDict((v, Counter(status_of(c['id']) for c in CASOS if c[key] == v)) for v in vals)
    by_prio = by('prioridad', ('Alta', 'Media', 'Baja'))
    by_tipo = by('tipo', ('Positivo', 'Negativo', 'Límite'))
    tecnicas = sorted({c['tecnica'] for c in CASOS}, key=lambda t: -sum(1 for c in CASOS if c['tecnica'] == t))
    by_tec = by('tecnica', tecnicas)
    by_canal = by('canal', ('UI', 'API'))
    return st, by_prio, by_tipo, by_tec, by_canal


def pct(n, d):
    return f'{(100 * n / d):.0f} %' if d else '—'


def table(title, d):
    L = [f'| {title} | Casos | PASSED | FAILED | DUDA | BLOCKED |', '|---|---|---|---|---|---|']
    for k, cnt in d.items():
        L.append(f'| {k} | {sum(cnt.values())} | {cnt["PASSED"]} | {cnt["FAILED"]} | {cnt["DUDA"]} | {cnt["BLOCKED"]} |')
    return L


# ---------------------------------------------------------------- 04 ejecución
def gen_exec():
    wb = Workbook()
    index_sheet(wb, f'Biyu – Entrega 2 – Ejecución de casos de prueba · Exportar CSV ({FECHA})')
    for c in CASOS:
        tc_sheet(wb, c)
    wb.save(OUT / '04-ejecucion-exp.xlsx')

    L = ['# Proyecto Biyu – Entrega 2 · Ejecución de casos de prueba (V2) · Exportar CSV (EXP)', '',
         'Planilla: `04-ejecucion-exp.xlsx` (una hoja por caso con "Reporte de ejecución 1", resultado obtenido, evidencia y defectos). '
         'Casos: [`casos/exp.md`](casos/exp.md) (CP-EXP-001 a CP-EXP-075, US-47). Reporte con métricas: [`06-reporte-de-ejecucion-exp.md`](06-reporte-de-ejecucion-exp.md). '
         'Defectos: [`05-reportes-de-defectos-exp.md`](05-reportes-de-defectos-exp.md).', '',
         '## Condiciones de la ejecución', '',
         f'- **Fecha:** {FECHA}.',
         f'- **Versión probada:** `main` en el commit `{COMMIT}` (código local, sin cambios). No se probó el deploy de Vercel.',
         f'- **Entorno:** {ENTORNO}.',
         f'- **Quién ejecutó:** **los 75 casos los ejecutó Claude Code (IA), supervisado por Santiago.** Ningún caso de EXP lo ejecutó el equipo a mano ni tenía un estado previo.',
         '- **Cómo se ejecutó:** un script de Playwright por caso ([`ejecucion/exp/`](ejecucion/exp/)). Cada caso crea su propio usuario (alta por API y entrada por la pantalla de ingreso, con la configuración inicial salteada), prepara los movimientos por la API (`create_transaction`, `delete_transaction`, `restore_transaction`, `upsert_fx_rate`, `delete_account`, categorías y cuentas) y recorre la pantalla Movimientos como lo haría una persona. Las descargas se leen del archivo que baja el navegador y se comparan byte a byte. Cada caso deja una captura del estado final en `evidencia/`.',
         '- **Datos:** usuarios `qa+cp-exp-<nnn>-001@example.com` (con un sufijo de corrida), contraseña `Clave123!` y montos ficticios (C14), en Supabase local. En producción no se creó nada.', '',
         '## Resultado por caso', '', '| ID | Prioridad | Historia | Status | Ejecutó | Defectos |', '|---|---|---|---|---|---|']
    L += [f'| {c["id"]} | {c["prioridad"]} | {c["trazas"]} | **{status_of(c["id"])}** | Claude Code (IA) | {defects_of(c["id"]) or "—"} |' for c in CASOS]
    L += ['', '## Detalle', '']
    for c in CASOS:
        r = RES[c['id']]
        L += [f'### {c["id"]} — {c["titulo"]} · **{status_of(c["id"])}**', '', '**Esperado:** ' + esperado(c), '',
              '**Obtenido:**', '', '```text', obtenido(c['id']), '```', '']
        if r.get('evidence'):
            L.append('**Evidencia:** ' + ' · '.join(f'[{Path(e).name}]({e})' for e in r['evidence']))
        if defects_of(c['id']):
            L.append(f'**Defectos:** {defects_of(c["id"])}')
        if r.get('note') and r['status'] != 'FAILED':
            L.append(f'**Notas:** {r["note"]}')
        L.append(f'**Ejecutó:** {EJECUTADO_POR}')
        L.append('')
    (OUT / '04-ejecucion-exp.md').write_text('\n'.join(L), encoding='utf-8', newline='\n')


# ---------------------------------------------------------------- 05 defectos
DET_LABEL = Font(name='Calibri', size=8, bold=True)
DET_VALUE = Font(name='Calibri', size=11)
DET_TITLE = Font(name='Calibri', size=11, bold=True)
DET_TOP = Alignment(wrap_text=True, vertical='top')
DET_CHARS_PER_LINE = 60


def det_height(text):
    lines = sum(max(1, -(-len(part) // DET_CHARS_PER_LINE)) for part in str(text).split('\n'))
    return max(15, 15 * lines)


def det_label(ws, row, text):
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill, c.border, c.alignment = DET_LABEL, LABEL, BOX, DET_TOP


def det_value(ws, row, text, font=DET_VALUE, border=False):
    c = ws.cell(row=row, column=2, value=text)
    c.font, c.alignment = font, DET_TOP
    if border:
        c.border = BOX
    ws.row_dimensions[row].height = det_height(text)


def det_notes(d):
    notes = [f'- Defect ID: {d["id"]}', f'- Estado: {d["estado"]}', f'- Issue de GitHub: {d.get("issue") or "sin issue"}',
             f'- Caso de prueba: {d["caso"]}', f'- Encontrado en: {d["encontrado"]}', f'- Reportó: {d["reporter"]}',
             f'- Evidencia: {d["evidencia"]}']
    if d.get('notas'):
        notes.append(f'- {d["notas"]}')
    return '\n'.join(notes)


def defect_sheet(wb, d):
    x = d['detalle']
    ws = wb.create_sheet(d['id'])
    ws.column_dimensions['A'].width = 31.83
    ws.column_dimensions['B'].width = 63.66
    row = 1
    det_label(ws, row, 'Title'); det_value(ws, row, f'[{d["id"]}] {d["titulo"]}', DET_TITLE, border=True); row += 2
    det_label(ws, row, 'Description / Overview'); det_value(ws, row, d['descripcion'], border=True); row += 2
    det_label(ws, row, 'Reproduction Steps')
    for i, paso in enumerate(d['pasos'], 1):
        det_value(ws, row, f'{i}- {paso}'); row += 1
    row += 1
    det_label(ws, row, 'Expected Behavior'); det_value(ws, row, d['esperado']); row += 2
    det_label(ws, row, 'Actual Behavior'); det_value(ws, row, d['obtenido']); row += 2
    det_label(ws, row, 'Incidence / Severity / Probability of reproduction')
    det_value(ws, row, f'Severidad {d["severidad"]} - {x["impacto"]}'); row += 1
    det_value(ws, row, x['reproduccion']); row += 2
    det_label(ws, row, 'Story and Acceptance Criteria affected'); det_value(ws, row, x['criterio']); row += 2
    det_label(ws, row, 'Browsers tested'); det_value(ws, row, x['navegadores']); row += 2
    det_label(ws, row, 'Environment'); det_value(ws, row, d['entorno']); row += 2
    det_label(ws, row, 'Notes'); det_value(ws, row, det_notes(d))


def gen_defects():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Index'
    heads = ['Defect ID', 'Title', 'Severity', 'Status', 'Test case', 'User story', 'GitHub issue']
    ws.append(['Biyu – Entrega 2 – Defect reports · Exportar CSV']); ws['A1'].font = Font(bold=True, size=13)
    ws.append(heads)
    for i in range(1, len(heads) + 1):
        c = ws.cell(row=2, column=i); c.fill = HEAD; c.font = WHITE; c.border = BOX
    for d in DEFECTOS:
        ws.append([d['id'], d['titulo'], d['severidad'], d['estado'], d['caso'], d['historia'], d.get('issue') or 'sin issue'])
        for i in range(1, len(heads) + 1):
            x = ws.cell(row=ws.max_row, column=i); x.border = BOX; x.alignment = WRAP
        ws.cell(row=ws.max_row, column=1).hyperlink = f"#'{d['id']}'!A1"
    for col, w in zip('ABCDEFG', (11, 60, 10, 34, 26, 30, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'
    for d in DEFECTOS:
        defect_sheet(wb, d)
    wb.save(OUT / '05-reportes-de-defectos-exp.xlsx')

    sev = Counter(d['severidad'] for d in DEFECTOS if not d['estado'].startswith('Cerrado'))
    L = ['# Proyecto Biyu – Entrega 2 · Reportes de defectos (V2) · Exportar CSV', '',
         'Planilla: `05-reportes-de-defectos-exp.xlsx` (un índice y una hoja por defecto, con el formato "Defect Report - Template" de la cátedra). '
         'La numeración sigue la de la Entrega 2: DEF-029 es el último de [`05-reportes-de-defectos.md`](05-reportes-de-defectos.md) y los de Exportar CSV empiezan en **DEF-030**. '
         'Todavía no tienen issue en GitHub: quedan listos para cargar.', '',
         'Escala de severidad y flujo de estados: `docs/07-plan-de-testing.md` §5. Severidad la fija quien reporta; **prioridad la fija el PO** (la que figura acá es la sugerida por quien reportó).', '',
         '## Resumen', '', '| ID | Título | Severidad | Prioridad sugerida | Estado | Caso | Historia |', '|---|---|---|---|---|---|---|']
    L += [f'| {d["id"]} | {d["titulo"]} | {d["severidad"]} | {d["prioridad"]} | {d["estado"]} | {d["caso"]} | {d["historia"]} |' for d in DEFECTOS]
    L += ['', '**Abiertos por severidad:** ' + ' · '.join(f'{k}: {sev.get(k, 0)}' for k in ('Crítica', 'Alta', 'Media', 'Baja')) + f' · Total abiertos: {sum(sev.values())}.', '']
    for d in DEFECTOS:
        L += [f'## {d["id"]} · {d["titulo"]}', '', '| Campo | Contenido |', '|---|---|',
              f'| Estado | {d["estado"]} |', f'| Severidad | {d["severidad"]} |', f'| Prioridad (sugerida) | {d["prioridad"]} |',
              f'| Encontrado en | {d["encontrado"]} |', f'| Caso de prueba | {d["caso"]} |', f'| Historia | {d["historia"]} |',
              f'| Issue | {d.get("issue") or "A crear"} |', f'| Reportó | {d["reporter"]} |', f'| Entorno | {d["entorno"]} |', '',
              d['descripcion'], '', '**Pasos para reproducir**', ''] + [f'{i}. {p}' for i, p in enumerate(d['pasos'], 1)] + [
              '', f'**Resultado esperado.** {d["esperado"]}', '', f'**Resultado obtenido.** {d["obtenido"]}', '',
              f'**Evidencia.** {d["evidencia"]}', '', f'**Notas.** {d["notas"]}', '']
    (OUT / '05-reportes-de-defectos-exp.md').write_text('\n'.join(L), encoding='utf-8', newline='\n')


# ---------------------------------------------------------------- 06 reporte
def gen_report():
    st, by_prio, by_tipo, by_tec, by_canal = metrics()
    total = len(CASOS)
    executed = st['PASSED'] + st['FAILED'] + st['DUDA']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte de ejecución - EXP E2'
    ws.append(['Biyu - entrega 2 - Reporte de ejecución de test cases · Exportar CSV'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells('A1:H1')
    heads = ['Test case ID', 'Prioridad', 'Título', 'User story', 'Camino feliz', 'Status', 'Ejecutó', 'Defectos']
    ws.append(heads)
    for i in range(1, 9):
        c = ws.cell(row=2, column=i); c.fill = HEAD; c.font = WHITE; c.border = BOX
    for c in CASOS:
        ws.append([c['id'], c['prioridad'], c['titulo'], c['trazas'], 'Sí' if c['feliz'] else 'No', status_of(c['id']), 'Claude Code (IA)', defects_of(c['id'])])
        for i in range(1, 9):
            x = ws.cell(row=ws.max_row, column=i); x.border = BOX; x.alignment = WRAP
        ws.cell(row=ws.max_row, column=6).fill = PatternFill('solid', fgColor=STATUS_FILL[status_of(c['id'])])
    for col, w in zip('ABCDEFGH', (14, 10, 58, 22, 12, 12, 18, 14)):
        ws.column_dimensions[col].width = w
    r0 = ws.max_row + 2
    ws.cell(row=r0, column=1, value='Resumen').font = BOLD
    summary = [('Planificados', total), ('Ejecutados (PASSED + FAILED + DUDA)', executed), ('PASSED', st['PASSED']), ('FAILED', st['FAILED']),
               ('DUDA', st['DUDA']), ('BLOCKED', st['BLOCKED']), ('No ejecutados', st['NO EJECUTADO']),
               ('% avance (ejecutados/planificados)', pct(executed, total)), ('% aprobados sobre ejecutados', pct(st['PASSED'], executed)),
               ('Ejecutados con IA', total), ('Ejecutados sin IA (a mano)', 0)]
    for i, (k, v) in enumerate(summary, r0 + 1):
        ws.cell(row=i, column=1, value=k).fill = LABEL
        ws.cell(row=i, column=2, value=v)
    r1 = r0 + len(summary) + 2
    ws.cell(row=r1, column=1, value='Por prioridad').font = BOLD
    for i, h in enumerate(['Prioridad', 'Casos', 'PASSED', 'FAILED', 'DUDA', 'BLOCKED'], 1):
        ws.cell(row=r1 + 1, column=i, value=h).fill = LABEL
    for j, (k, cnt) in enumerate(by_prio.items(), r1 + 2):
        for i, v in enumerate([k, sum(cnt.values()), cnt['PASSED'], cnt['FAILED'], cnt['DUDA'], cnt['BLOCKED']], 1):
            ws.cell(row=j, column=i, value=v)
    wb.save(OUT / '06-reporte-de-ejecucion-exp.xlsx')

    alta = [c for c in CASOS if c['prioridad'] == 'Alta']
    alta_no = [c['id'] for c in alta if status_of(c['id']) == 'NO EJECUTADO']
    sustituidas = ['CP-EXP-030', 'CP-EXP-049', 'CP-EXP-050', 'CP-EXP-057', 'CP-EXP-058', 'CP-EXP-059', 'CP-EXP-064', 'CP-EXP-068',
                   'CP-EXP-069', 'CP-EXP-070', 'CP-EXP-071', 'CP-EXP-073', 'CP-EXP-074']
    L = ['# Proyecto Biyu – Entrega 2 · Reporte de ejecución de casos de prueba (V2) · Exportar CSV (EXP)', '',
         'Planilla: `06-reporte-de-ejecucion-exp.xlsx` (hoja "Reporte de ejecución - EXP E2", mismo formato que el de la Entrega 1: ID, prioridad, título, historia, camino feliz, status y defectos, más el resumen). Formato del contenido: `docs/07-plan-de-testing.md` §6. '
         'Registro por caso en [`04-ejecucion-exp.md`](04-ejecucion-exp.md); defectos en [`05-reportes-de-defectos-exp.md`](05-reportes-de-defectos-exp.md).', '',
         '| Dato | Valor |', '|---|---|',
         f'| Versión probada | `main` en el commit `{COMMIT}` (código local, sin cambios). No se probó el deploy de Vercel |',
         '| Entorno | Local: Vite en `http://127.0.0.1:5173` + Supabase local. Chromium (Playwright 1.63), es-AR, America/Argentina/Buenos_Aires, reloj del navegador fijado en 09/10/2026 12:00 |',
         f'| Fecha | {FECHA} |',
         f'| Ejecutó | {EJECUTADO_POR}. Los 75 casos, con los scripts de `ejecucion/exp/` (`casos-a.mjs`, `casos-b.mjs` y `casos-c.mjs`) |',
         '| Evidencia | `ejecucion/exp/resultados-a.json`, `resultados-b.json` y `resultados-c.json` (resultado y observaciones de cada caso) y 75 capturas en `evidencia/` (`CP-EXP-nnn.jpg`) |', '',
         '## 1. Planificados, ejecutados, pasados, fallados y bloqueados', '',
         'Status vigente de cada caso: el de su única ejecución (no hubo re-test).', '',
         '| Métrica | Casos | % |', '|---|---|---|',
         f'| Planificados | **{total}** | 100 % |', f'| Ejecutados (PASSED + FAILED + DUDA) | **{executed}** | {pct(executed, total)} de avance |',
         f'| PASSED | **{st["PASSED"]}** | {pct(st["PASSED"], executed)} de los ejecutados |', f'| FAILED | **{st["FAILED"]}** | {pct(st["FAILED"], executed)} de los ejecutados |',
         f'| DUDA (pasa, con dudas) | **{st["DUDA"]}** | {pct(st["DUDA"], executed)} de los ejecutados |', f'| BLOCKED | **{st["BLOCKED"]}** | — |', f'| No ejecutados | {st["NO EJECUTADO"]} | — |', '',
         '- **FAILED:** CP-EXP-062 (DEF-030), CP-EXP-063 y CP-EXP-065 (DEF-031) y CP-EXP-064 (DEF-032). Los cuatro se reprodujeron más de una vez y el defecto es de la aplicación, no del caso.',
         '- **DUDA:** CP-EXP-030 y CP-EXP-071. Los pasos finales piden importar el CSV en Google Sheets; no se pudo usar, así que el oráculo fue un lector RFC 4180 propio (12 columnas, una fila y descripción multilínea en una sola celda). Repetir con Google Sheets para cerrarlos.',
         '- **PASSED sin valor hoy:** CP-EXP-070. Exige que tocar fuera no cierre la hoja durante «Exportando…», pero hoy tocar fuera nunca la cierra (DEF-030): pasaría aunque el bloqueo no existiera. Repetir cuando se corrija DEF-030.',
         '- **Lo que sí se cumple:** el contenido del CSV pasó todas las comprobaciones de formato (BOM, doce columnas, CRLF), montos y tipo de cambio, escape de comas, comillas y saltos de línea, neutralización de fórmulas, orden, aislamiento entre usuarios (UI y API) y volumen (1.000, 1.500 y 2.000 filas sin faltantes).', '',
         '## 2. Cobertura por prioridad, tipo, técnica y canal', ''] + table('Prioridad', by_prio) + [''] + table('Tipo', by_tipo) + [''] + table('Técnica', by_tec) + [''] + table('Canal', by_canal) + ['',
         f'Camino feliz: {sum(1 for c in CASOS if c["feliz"])} casos. Negativos y de límite: {sum(1 for c in CASOS if not c["feliz"])} casos.', '',
         '**Pruebas automatizadas existentes:** el test unitario del dominio de exportación (`tests/domain/exportCsv.test.ts`) pasó 18/18 en esta corrida. No se corrieron el resto de Vitest, pgTAP ni Playwright smoke.', '',
         '## 3. Quién ejecutó cada caso: con IA y sin IA', '',
         '| Forma de ejecución | Casos | Cuáles |', '|---|---:|---|',
         f'| **Con IA** (Claude Code, supervisado por Santiago) | **{total}** | CP-EXP-001 a CP-EXP-075 |', '| Sin IA (a mano por el equipo) | 0 | — |', '',
         'El equipo no tenía resultados previos de EXP, así que todo el módulo se ejecutó con IA. En estos casos la IA tuvo que sustituir la herramienta o el oráculo del caso (detalle en el punto 7): ' + ', '.join(sustituidas) + '.', '',
         '## 4. Trazabilidad', '',
         'Los 75 casos cubren los criterios CA-1 a CA-25, el requisito de cambio de URL, la completitud y el estado Error de US-47 (matriz en `casos/exp.md` §1). Todos tienen al menos un caso ejecutado. Los casos que fallaron cubren CA-23 (CP-EXP-062) y CA-24 (CP-EXP-063 y 064), más el requisito «todo cambio de URL» (CP-EXP-065).', '',
         '**Invariantes y restricciones:** C1, C2, C5, C7, C9, C10, C11 e I10 tienen casos ejecutados y en verde. Las demás invariantes de V2 se informan en otros módulos.', '',
         '## 5. Defectos por severidad', '', 'Detalle completo en `05-reportes-de-defectos-exp`.', '',
         '| Severidad | Abiertos | IDs |', '|---|---|---|', '| Crítica | 0 | — |', '| Alta | 0 | — |', '| Media | 3 | DEF-030, DEF-031, DEF-032 |', '| Baja | 0 | — |', '| **Total abiertos** | **3** | DEF-030, DEF-031, DEF-032 |', '',
         'Los tres defectos son de la hoja «Exportar movimientos» (cómo se cierra); ninguno afecta el contenido del CSV. Ninguno tiene issue todavía.', '',
         '## 6. Criterios de salida (docs/07-plan-de-testing.md §6)', '',
         '| Criterio | Estado |', '|---|---|',
         '| 1. Cero defectos de severidad Crítica abiertos | ✅ Cumplido |',
         '| 2. Cero defectos de severidad Alta abiertos, o justificados por el PO | ✅ Cumplido |',
         f'| 3. Todos los casos de prioridad Alta ejecutados | {"✅ Cumplido" if not alta_no else "❌ No cumplido"}: {len(alta) - len(alta_no)} de {len(alta)} (con FAILED: CP-EXP-064; con DUDA: CP-EXP-030 y CP-EXP-071) |',
         '| 4. Toda invariante del alcance con al menos un caso ejecutado | Parcial: las de EXP sí (C1, C2, C5, C7, C9, C10, C11, I10); las demás de V2 van en otros módulos |', '',
         '**Conclusión:** EXP cumple los criterios 1, 2 y 3 y no tiene casos sin ejecutar ni bloqueados. Hay tres defectos Medios abiertos (DEF-030 a DEF-032) sobre el cierre de la hoja, que hay que corregir y volver a probar junto con CP-EXP-062 a 065 y CP-EXP-070. Quedan por repetir con Google Sheets CP-EXP-030 y CP-EXP-071.', '',
         '## 7. Desvíos respecto de lo planificado', '',
         '1. **Ejecución asistida por un script.** El plan prevé ejecución manual por un integrante que no diseñó el caso. Acá los pasos los ejecutó Claude Code con Playwright, siguiendo cada caso como lo haría una persona y leyendo el archivo descargado como oráculo. La propiedad cruzada del plan (`docs/07-plan-de-testing.md` §1) no se aplicó: conviene que cada integrante revise la evidencia de su módulo.',
         '2. **Entorno local, no el deploy.** Todo se ejecutó contra el código local y Supabase local; los resultados pueden diferir del deploy de Vercel si este está desactualizado. Solo Chromium: no se probó Firefox ni Safari.',
         '3. **Reloj fijado.** El reloj del navegador se fijó en 09/10/2026 12:00 (America/Argentina/Buenos_Aires), como piden los casos CP-EXP-002 y 004 a 006. El reloj del servidor no se tocó.',
         '4. **Preparación por API.** Las transacciones, el catálogo y las referencias de tipo de cambio se cargaron por API (EXP-PREP-01 a 06), no por las pantallas. Los casos 040 a 042 y 068 sembraron 1.000, 1.500 y 2.000 transacciones con un script de Node (8 pedidos en paralelo) en lugar del script de PowerShell 7 de EXP-PREP-07, que no está instalado.',
         '5. **Herramientas sustituidas.** (a) Red lenta (049, 050, 064, 069 y 070): demora fija de 4 s (3 s en el 064) en la lectura de `/rest/v1/transactions`, en lugar de «1 kbit/s, 2000 ms» de DevTools. (b) Sin conexión (057 a 059): corte de red del contexto de Playwright. (c) HTTP 500 y monto inválido (073 y 074): ruta interceptada de Playwright en lugar del código pegado en la consola (EXP-PREP-08 y 09), con el mismo efecto sobre el `fetch`. (d) CP-EXP-068: el alta concurrente se hizo por API entre la primera y la segunda página de la lectura, no desde una segunda ventana de otro navegador. (e) CP-EXP-030 y 071: lector RFC 4180 propio en lugar de Google Sheets. (f) 062 y 070: clic fuera de la hoja en la posición (5, 5) de una ventana de 420 × 900.',
         '6. **Cuotas (023 y 024).** La preparación pide PR-07; se creó la compra con `create_transaction` en una cuenta «Visa BBVA» de tipo tarjeta de crédito, porque la base solo admite cuotas con tarjeta de crédito.',
         '7. **Bajas.** `delete_transaction` responde 204 sin cuerpo (ver DEF-029); el script lo toma como éxito en 017, 019, 056 y los demás casos que eliminan una transacción.',
         '8. **Casos reescritos por errores del script, no del caso ni de la app.** Se rehicieron CP-EXP-017 (la baja responde 204), CP-EXP-023 y 024 (faltaba la tarjeta de crédito) y CP-EXP-062 a 066 (esperas y verificaciones insuficientes en la primera versión; CP-EXP-066 se ejecutó tecleando la URL de destino, como pide el caso). No se cambió el resultado de ningún caso a mano.',
         '9. **Dependencias.** Para correr el import de Excel hubo que instalar `read-excel-file` y `write-excel-file` (estaban en `package.json` pero no en `node_modules`); no afecta a EXP.',
         '10. **Aviso de suscripciones en las capturas.** En varias capturas se ve «No pudimos cargar tus suscripciones vencidas»: la Edge Function de la puesta al día de suscripciones no estaba levantada en el Supabase local. No afecta a Exportar CSV (no se usa) y no cambia ningún resultado.',
         '11. **Evidencia.** Cada caso guarda una captura del estado final de la pantalla (`evidencia/CP-EXP-nnn.jpg`) y las observaciones del script (nombre y tamaño del archivo, fila exportada, mensajes). El contenido de los CSV no se guardó en el repositorio.', '',
         '## 8. Datos de prueba', '',
         'Todo en Supabase **local**, con usuarios `qa+cp-exp-<nnn>-001@example.com` (los de aislamiento, `-A` y `-B`) y montos ficticios. `supabase db reset` los borra. En producción no se creó nada. La contraseña de prueba (`Clave123!`) está en `ejecucion/exp/lib.mjs` y solo existe en la base local.', '',
         '## Cómo se reproduce', '',
         '```bash',
         'supabase start && supabase migration up',
         'VITE_SUPABASE_URL=http://127.0.0.1:54321 VITE_SUPABASE_ANON_KEY=<anon key local> npx vite --port 5173 --strictPort',
         'node entrega-2/ejecucion/exp/casos-a.mjs   # CP-EXP-001 a 039 (menos 017, 023 y 024)',
         'node entrega-2/ejecucion/exp/casos-b.mjs   # 017, 023, 024 y 040 a 075 (menos 062 a 066)',
         'node entrega-2/ejecucion/exp/casos-c.mjs   # 062 a 066',
         'PYTHONUTF8=1 python entrega-2/ejecucion/exp/generar_exp.py',
         '```', '',
         '- Cada script escribe `resultados-<a|b|c>.json` y las capturas de `evidencia/`. Solo apuntan al stack local: nunca a producción.',
         '- `generar_exp.py` necesita `openpyxl` y arma los tres `.md` y las tres planillas. Los resultados y el texto de lo obtenido salen de `resultados-*.json` y `obtenido.json`; los defectos, de `defectos.json`.', '']
    (OUT / '06-reporte-de-ejecucion-exp.md').write_text('\n'.join(L), encoding='utf-8', newline='\n')
    return st


if __name__ == '__main__':
    gen_exec()
    gen_defects()
    st = gen_report()
    print(json.dumps(dict(st), ensure_ascii=False))

"""Genera los entregables de la Parte B (Markdown + planillas .xlsx con el formato de la cátedra)
desde tres fuentes: casos.json (exportado de casos.mjs), ../ejecucion/resultados.json y defectos.json.

Uso (desde la raíz del repo):
  node -e "import('./entrega-1/fuentes/casos.mjs').then(m=>process.stdout.write(JSON.stringify(m.CASOS)))" > entrega-1/fuentes/casos.json
  python3 entrega-1/fuentes/generar.py
  python3 entrega-1/fuentes/generar.py defectos   (solo 04: .md y .xlsx de defectos)
Necesita openpyxl. La planilla de defectos también lee defectos-detalle.json.
"""
import json
import sys
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent
CASOS = json.loads((ROOT / 'fuentes' / 'casos.json').read_text())
RUN = json.loads((ROOT / 'ejecucion' / 'resultados.json').read_text())
DEFECTOS = json.loads((ROOT / 'fuentes' / 'defectos.json').read_text())
RES = {r['id']: r for r in RUN['results']}
RETEST = {r['id']: r for r in RUN['retests']}
# Ejecución 2: re-test parcial (solo los casos que se volvieron a correr). El status vigente de
# cada caso es el de su última ejecución.
EJ2 = RUN.get('ejecucion2') or {'results': []}
RES2 = {r['id']: r for r in EJ2['results']}
FECHA2 = EJ2.get('today', '')
COMMIT2 = EJ2.get('commit', '')
BY_ID = {c['id']: c for c in CASOS}

PROYECTO = 'Biyu'
CREADO_POR = 'Equipo Biyu (skill /new-test-case, revisado por spec-critic)'
EJECUTADO_POR = 'Claude Code con el runner entrega-1/ejecucion/run.mjs, supervisado por Joaquin Nuñez'
FECHA = RUN['today']
COMMIT = RUN.get('commit') or '44f1519'
ENTORNO = f'Local: Vite (http://localhost:5180) + Supabase local · main {COMMIT} · Chromium (Playwright) 390×844 móvil'
ENTORNO2 = f'Local: Vite (http://localhost:5180) + Supabase local · main {COMMIT2} · Chromium (Playwright) 390×844 móvil'
MODULOS = OrderedDict([('ACC', 'Acceso y autorización'), ('CFG', 'Configuración'), ('REG', 'Registro y baja'),
                       ('CUO', 'Cuotas'), ('MON', 'Monedas'), ('DAS', 'Dashboard')])

LABEL = PatternFill('solid', fgColor='FFFFFFCC')
EXEC = PatternFill('solid', fgColor='FFC6EFCE')
HEAD = PatternFill('solid', fgColor='FF1B4D3E')
STATUS_FILL = {'PASSED': 'FFC6EFCE', 'FAILED': 'FFFFC7CE', 'BLOCKED': 'FFFFEB9C', 'NO EJECUTADO': 'FFD9D9D9'}
THIN = Side(style='thin', color='FFBFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
BOLD = Font(bold=True)
WHITE = Font(bold=True, color='FFFFFFFF')


def last_run(cid):
    return RES2.get(cid) or RES.get(cid)


def status_of(cid):
    r = last_run(cid)
    return r['status'] if r else 'NO EJECUTADO'


def status_1(cid):
    r = RES.get(cid)
    return r['status'] if r else 'NO EJECUTADO'


def defects_of(cid):
    """Defectos de la última ejecución; si ya no están, se nombran los que había en la ejecución 1."""
    now = (last_run(cid) or {}).get('defects', [])
    before = RES.get(cid, {}).get('defects', [])
    if now or not before or cid not in RES2:
        return ', '.join(now)
    return f'{", ".join(before)} (corregido)'


def text_height(text, chars_per_line):
    """Alto de fila (puntos) para que un texto con ajuste de línea se lea entero."""
    lines = sum(max(1, -(-len(part) // chars_per_line)) for part in str(text).split('\n'))
    return max(15, 15 * lines + 2)


def cell(ws, ref, value, fill=None, bold=False, merge=None):
    ws[ref] = value
    ws[ref].alignment = WRAP
    ws[ref].border = BOX
    if fill:
        ws[ref].fill = fill
    if bold:
        ws[ref].font = BOLD
    if merge:
        ws.merge_cells(merge)
    return ws[ref]


def tc_sheet(wb, c, with_exec):
    """Una hoja por caso, con la misma grilla que 'TC - Template' / las hojas TC-xx del ejemplo."""
    ws = wb.create_sheet(c['id'].replace('CP-', ''))
    for col, w in zip('ABCDEFG', (22, 46, 22, 46, 22, 40, 40)):
        ws.column_dimensions[col].width = w
    cell(ws, 'A1', 'Test Case ID', LABEL, True); cell(ws, 'B1', c['id'])
    cell(ws, 'C1', 'Descripción / Funcionalidad', LABEL, True); cell(ws, 'D1', c['funcionalidad'], merge='D1:E1')
    cell(ws, 'A2', 'Creado por', LABEL, True); cell(ws, 'B2', CREADO_POR)
    cell(ws, 'C2', 'Versión', LABEL, True); cell(ws, 'D2', '2 (nuevo en la Entrega 1)' if c.get('nuevo') else '2', merge='D2:E2')
    cell(ws, 'A3', 'Reporte de ejecución 1', EXEC, True, merge='A3:B3')
    cell(ws, 'C3', 'Reporte de ejecución 2', EXEC, True); cell(ws, 'D3', 'Reporte de ejecución 3', EXEC, True, merge='D3:E3')
    if with_exec:
        st = status_1(c['id'])
        cell(ws, 'A4', f'{st} · {FECHA}', PatternFill('solid', fgColor=STATUS_FILL[st]), True, merge='A4:B4')
        if c['id'] in RES2:
            st2 = RES2[c['id']]['status']
            cell(ws, 'C4', f'{st2} · {FECHA2}', PatternFill('solid', fgColor=STATUS_FILL[st2]), True)
    cell(ws, 'A5', 'Prioridad (Alta, Media, Baja)', LABEL, True); cell(ws, 'B5', c['prioridad'])
    cell(ws, 'C5', 'Camino feliz', LABEL, True); cell(ws, 'D5', 'Sí' if c['feliz'] else 'No')
    cell(ws, 'A6', 'Técnica / Tipo', LABEL, True); cell(ws, 'B6', f"{c['tecnica']} · {c['tipo']}")
    cell(ws, 'C6', 'Trazabilidad', LABEL, True); cell(ws, 'D6', c['trazas'], merge='D6:E6')
    row = 8
    cell(ws, f'A{row}', 'S #', LABEL, True); cell(ws, f'B{row}', 'Pre-requisitos', LABEL, True)
    cell(ws, f'C{row}', 'S #', LABEL, True); cell(ws, f'D{row}', 'Datos de prueba', LABEL, True, merge=f'D{row}:E{row}')
    n = max(len(c['pre']), len(c['datos']), 1)
    for i in range(n):
        r = row + 1 + i
        if i < len(c['pre']):
            cell(ws, f'A{r}', i + 1); cell(ws, f'B{r}', c['pre'][i])
        if i < len(c['datos']):
            cell(ws, f'C{r}', i + 1); cell(ws, f'D{r}', f'{c["datos"][i][0]}:'); cell(ws, f'E{r}', c['datos'][i][1])
        pre = c['pre'][i] if i < len(c['pre']) else ''
        dato = c['datos'][i][1] if i < len(c['datos']) else ''
        ws.row_dimensions[r].height = max(text_height(pre, 42), text_height(dato, 20))
    row += n + 2
    cell(ws, f'A{row}', 'Título del CP', LABEL, True); cell(ws, f'B{row}', c['titulo'], merge=f'B{row}:E{row}')
    row += 2
    cell(ws, f'A{row}', 'Historias de Usuario', LABEL, True); cell(ws, f'B{row}', ', '.join(c['historias']))
    cell(ws, f'C{row}', 'Estado Automatización', LABEL, True)
    auto = {'Sí': 'Automatizable (candidato)', 'No': 'No automatizable', 'V3': 'Candidato a V3'}[c['automatizable']]
    cell(ws, f'D{row}', f"{auto}. Ejecución de esta entrega asistida por runner Playwright.", merge=f'D{row}:E{row}')
    row += 2
    cell(ws, f'A{row}', 'Pasos #', LABEL, True); cell(ws, f'B{row}', 'Detalle de pasos', LABEL, True, merge=f'B{row}:C{row}')
    cell(ws, f'D{row}', 'Resultado Esperado', LABEL, True, merge=f'D{row}:E{row}')
    for i, (paso, esperado) in enumerate(c['pasos'], 1):
        r = row + i
        cell(ws, f'A{r}', i); cell(ws, f'B{r}', paso, merge=f'B{r}:C{r}'); cell(ws, f'D{r}', esperado, merge=f'D{r}:E{r}')
        # Excel no ajusta solo el alto de una fila con celdas combinadas: se calcula acá.
        ws.row_dimensions[r].height = max(text_height(paso, 62), text_height(esperado, 62))
    if with_exec:
        row += len(c['pasos']) + 2
        runs = [('', RES.get(c['id'], {}), ENTORNO)]
        if c['id'] in RES2:
            runs = [(' · ejecución 1', RES.get(c['id'], {}), ENTORNO), (' · ejecución 2', RES2[c['id']], ENTORNO2)]
        for suffix, r, entorno in runs:
            cell(ws, f'A{row}', f'Resultado obtenido{suffix}', EXEC, True); cell(ws, f'B{row}', r.get('obtained', 'No ejecutado'), merge=f'B{row}:E{row}')
            ws.row_dimensions[row].height = max(30, 15 * (r.get('obtained', '').count('\n') + 2))
            cell(ws, f'A{row + 1}', 'Evidencia', EXEC, True); cell(ws, f'B{row + 1}', '\n'.join(r.get('evidence', [])) or '—', merge=f'B{row + 1}:E{row + 1}')
            cell(ws, f'A{row + 2}', 'Defectos', EXEC, True); cell(ws, f'B{row + 2}', ', '.join(r.get('defects', [])) or '—', merge=f'B{row + 2}:E{row + 2}')
            cell(ws, f'A{row + 3}', 'Notas del ejecutor', EXEC, True); cell(ws, f'B{row + 3}', r.get('notes') or '—', merge=f'B{row + 3}:E{row + 3}')
            cell(ws, f'A{row + 4}', 'Entorno / ejecutó', EXEC, True); cell(ws, f'B{row + 4}', f'{entorno}\n{EJECUTADO_POR}', merge=f'B{row + 4}:E{row + 4}')
            row += 6
    return ws


def index_sheet(wb, title, with_exec):
    ws = wb.active
    ws.title = 'Índice'
    heads = ['Test case ID', 'Prioridad', 'Título', 'User story', 'Camino feliz', 'Tipo', 'Técnica']
    if with_exec:
        heads += ['Status', 'Defectos']
    ws.append([title]); ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
    ws.append(heads)
    for i, h in enumerate(heads, 1):
        c = ws.cell(row=2, column=i); c.fill = HEAD; c.font = WHITE; c.border = BOX; c.alignment = WRAP
    for c in CASOS:
        row = [c['id'], c['prioridad'], c['titulo'], ', '.join(c['historias']), 'Sí' if c['feliz'] else 'No', c['tipo'], c['tecnica']]
        if with_exec:
            row += [status_of(c['id']), defects_of(c['id'])]
        ws.append(row)
        for i in range(1, len(row) + 1):
            x = ws.cell(row=ws.max_row, column=i); x.border = BOX; x.alignment = WRAP
        if with_exec:
            ws.cell(row=ws.max_row, column=8).fill = PatternFill('solid', fgColor=STATUS_FILL[status_of(c['id'])])
    for col, w in zip('ABCDEFGHI', (14, 10, 58, 18, 12, 11, 22, 12, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'
    return ws


# ---------------------------------------------------------------- métricas
def metrics():
    st = Counter(status_of(c['id']) for c in CASOS)
    by_mod = OrderedDict()
    for m in MODULOS:
        cs = [c for c in CASOS if c['id'].startswith(f'CP-{m}-')]
        by_mod[m] = Counter(status_of(c['id']) for c in cs)
        by_mod[m]['Total'] = len(cs)
    by_prio = OrderedDict((p, Counter(status_of(c['id']) for c in CASOS if c['prioridad'] == p)) for p in ('Alta', 'Media', 'Baja'))
    by_tipo = Counter(c['tipo'] for c in CASOS)
    by_tec = Counter(c['tecnica'] for c in CASOS)
    return st, by_mod, by_prio, by_tipo, by_tec


def pct(n, d):
    return f'{(100 * n / d):.0f} %' if d else '—'


# ---------------------------------------------------------------- 02 especificación
def gen_spec():
    wb = Workbook()
    index_sheet(wb, 'Biyu – Entrega 1 – Especificación de casos de prueba (V1)', False)
    for c in CASOS:
        tc_sheet(wb, c, False)
    wb.save(ROOT / '02-especificacion-casos-de-prueba.xlsx')

    _, by_mod, by_prio, by_tipo, by_tec = metrics()
    L = ['# Proyecto Biyu – Entrega 1 · Especificación de casos de prueba (V1)', '',
         f'Planilla con una hoja por caso (formato de la plantilla de la cátedra): `02-especificacion-casos-de-prueba.xlsx`. '
         f'Fuente única: `fuentes/casos.mjs`.', '',
         '## Qué contiene', '',
         f'- **{len(CASOS)} casos**: los 68 del catálogo de V1 (`docs/10-catalogo-casos-v1.md`, mismos IDs y oráculos) '
         f'más **6 casos nuevos** de esta entrega para las historias que no tenían ninguno: CP-ACC-012 (US-66) y CP-CFG-011 a CP-CFG-015 (US-68).',
         '- Cada caso declara técnica de diseño, tipo, prioridad, si es **camino feliz**, precondiciones, datos de prueba y pasos con su resultado esperado.',
         '- Los negativos llevan una **variante API**: el mismo ataque directo contra Supabase, salteando la UI (C6: la validación real es la de Postgres).',
         '- IDs: `CP-<módulo>-<nnn>`. Módulos: ' + ' · '.join(f'`{k}` {v}' for k, v in MODULOS.items()) + '.', '',
         '## Cobertura del diseño', '',
         '| Módulo | Casos | Alta | Media | Baja |', '|---|---|---|---|---|']
    for m, name in MODULOS.items():
        cs = [c for c in CASOS if c['id'].startswith(f'CP-{m}-')]
        cnt = Counter(c['prioridad'] for c in cs)
        L.append(f'| {m} · {name} | {len(cs)} | {cnt["Alta"]} | {cnt["Media"]} | {cnt["Baja"]} |')
    cnt = Counter(c['prioridad'] for c in CASOS)
    L += [f'| **Total** | **{len(CASOS)}** | **{cnt["Alta"]}** | **{cnt["Media"]}** | **{cnt["Baja"]}** |', '',
          '| Tipo | Casos |', '|---|---|'] + [f'| {k} | {v} |' for k, v in by_tipo.most_common()] + ['',
          '| Técnica | Casos |', '|---|---|'] + [f'| {k} | {v} |' for k, v in by_tec.most_common()] + ['',
          f'Casos de camino feliz: {sum(1 for c in CASOS if c["feliz"])}.', '',
          '## Índice', '', '| ID | Prioridad | Título | Historia | Camino feliz | Tipo |', '|---|---|---|---|---|---|']
    L += [f'| {c["id"]}{" *(nuevo)*" if c.get("nuevo") else ""} | {c["prioridad"]} | {c["titulo"]} | {", ".join(c["historias"])} | {"Sí" if c["feliz"] else "No"} | {c["tipo"]} |' for c in CASOS]
    L += ['', '## Casos', '']
    for m, name in MODULOS.items():
        L += [f'### {m} · {name}', '']
        for c in [c for c in CASOS if c['id'].startswith(f'CP-{m}-')]:
            L += [f'#### {c["id"]} — {c["titulo"]}{" *(nuevo)*" if c.get("nuevo") else ""}', '',
                  f'| Campo | Contenido |', '|---|---|',
                  f'| Funcionalidad | {c["funcionalidad"]} |',
                  f'| Historias de usuario | {", ".join(c["historias"])} |',
                  f'| Trazabilidad | {c["trazas"]} |',
                  f'| Técnica · Tipo | {c["tecnica"]} · {c["tipo"]} |',
                  f'| Prioridad · Camino feliz | {c["prioridad"]} · {"Sí" if c["feliz"] else "No"} |',
                  f'| Automatización | {c["automatizable"]} |',
                  f'| Pre-requisitos | {"<br>".join(c["pre"])} |',
                  f'| Datos de prueba | {"<br>".join(f"{k}: {v}" for k, v in c["datos"]) or "—"} |', '',
                  '| # | Paso | Resultado esperado |', '|---|---|---|']
            L += [f'| {i} | {p} | {e} |' for i, (p, e) in enumerate(c['pasos'], 1)]
            L.append('')
    L += ['## Supuestos', '',
          '1. Los 68 casos conservan el oráculo del catálogo del repo; donde el rediseño (ADR-023/024) cambió el camino en pantalla se reescribió el paso, no el resultado esperado.',
          '2. "Creado por": el catálogo se diseñó con la skill `/new-test-case` a partir de la spec y lo revisó el agente `spec-critic` (`docs/10-catalogo-casos-v1.md` §Revisión). No hay registro de qué integrante diseñó cada caso, así que no se asigna uno por caso.',
          '3. Los 6 casos nuevos (US-66, US-68) se diseñaron en esta entrega desde los criterios de aceptación de sus issues.', '']
    (ROOT / '02-especificacion-casos-de-prueba.md').write_text('\n'.join(L))


# ---------------------------------------------------------------- 03 ejecución
def gen_exec():
    wb = Workbook()
    title = f'Biyu – Entrega 1 – Ejecución de casos de prueba ({FECHA}' + (f', re-test {FECHA2})' if RES2 else ')')
    index_sheet(wb, title, True)
    for c in CASOS:
        tc_sheet(wb, c, True)
    wb.save(ROOT / '03-ejecucion-casos-de-prueba.xlsx')

    L = ['# Proyecto Biyu – Entrega 1 · Ejecución de casos de prueba (V1)', '',
         f'Planilla: `03-ejecucion-casos-de-prueba.xlsx` (una hoja por caso con "Reporte de ejecución 1", resultado obtenido, evidencia y defectos; '
         f'los casos re-testeados tienen además "Reporte de ejecución 2").', '',
         '## Condiciones de la ejecución', '',
         f'- **Fecha:** {FECHA} (hoy según el reloj de la corrida, hora argentina).',
         f'- **Versión probada:** `main` en el commit `{COMMIT}` (incluye los fixes de DEF-004 y DEF-001 mergeados ese día).',
         f'- **Entorno:** {ENTORNO}.',
         f'- **Cómo se ejecutó:** {EJECUTADO_POR}. El runner recorre cada caso por la UI como lo haría una persona (Playwright, emulación de celular), ejecuta las variantes API con la anon key y sesiones reales de usuarios de prueba, y usa consultas directas a la base local solo como oráculo. Cada caso guarda su resultado obtenido y, si aplica, una captura en `evidencia/`.',
         f'- **Datos:** usuarios y montos ficticios creados por la corrida (`{RUN["run"]}`) en Supabase local. En producción no se creó nada.',
         '- **Además:** Vitest 239/239 y pgTAP 158/158 en verde sobre el mismo commit (ver `05-reporte-de-ejecucion`).', '']
    if RES2:
        L += ['## Ejecución 2 (re-test)', '',
              f'- **Fecha:** {FECHA2}. **Versión probada:** `main` en el commit `{COMMIT2}`. **Corrida:** `{EJ2["run"]}`, mismo entorno y runner.',
              f'- **Qué se re-ejecutó:** {", ".join(RES2)}. {EJ2["motivo"]}',
              '- **Status vigente:** el de la última ejecución de cada caso. La ejecución 1 queda registrada en la hoja de cada caso.', '']
    L += ['## Resultado por caso', '', '| ID | Prioridad | Historia | Status | Defectos |', '|---|---|---|---|---|']
    L += [f'| {c["id"]} | {c["prioridad"]} | {", ".join(c["historias"])} | **{status_of(c["id"])}**{" (ejecución 2)" if c["id"] in RES2 else ""} | {defects_of(c["id"]) or "—"} |' for c in CASOS]
    L += ['', '## Detalle', '']
    for c in CASOS:
        L += [f'### {c["id"]} — {c["titulo"]} · **{status_of(c["id"])}**', '',
              f'**Esperado:** ' + ' '.join(e for _, e in c['pasos']), '']
        runs = [('', RES.get(c['id'], {}))]
        if c['id'] in RES2:
            runs = [(f' (ejecución 1, {FECHA}, {status_1(c["id"])})', RES.get(c['id'], {})),
                    (f' (ejecución 2, {FECHA2}, {RES2[c["id"]]["status"]})', RES2[c['id']])]
        for suffix, r in runs:
            L += [f'**Obtenido{suffix}:**', '', '```text', r.get('obtained', 'No ejecutado'), '```', '']
            if r.get('evidence'):
                L.append('**Evidencia:** ' + ' · '.join(f'[{Path(e).name}]({e})' for e in r['evidence']))
            if r.get('defects'):
                L.append(f'**Defectos:** {", ".join(r["defects"])}')
            if r.get('notes'):
                L.append(f'**Notas:** {r["notes"]}')
            L.append('')
    (ROOT / '03-ejecucion-casos-de-prueba.md').write_text('\n'.join(L))


# ---------------------------------------------------------------- 04 defectos
# Layout de la hoja "Defect Report - Template" de la cátedra (TPO - TC and Defect Template.xlsx):
# etiqueta en A (Calibri 8, negrita, fondo amarillo), valor en B, una fila en blanco entre
# secciones y un paso por fila. Lo que el template no tiene (ID, estado, issue, caso, evidencia,
# corrección) va en Notes. La prioridad no se completa: la fija el PO.
DET_LABEL = Font(name='Calibri', size=8, bold=True)
DET_VALUE = Font(name='Calibri', size=11)
DET_TITLE = Font(name='Calibri', size=11, bold=True)
DET_TOP = Alignment(wrap_text=True, vertical='top')
DET_CHARS_PER_LINE = 60  # columna B de 63,66 de ancho, Calibri 11


def det_height(text):
    lines = sum(max(1, -(-len(part) // DET_CHARS_PER_LINE)) for part in str(text).split('\n'))
    return max(15, 15 * lines)


def det_label(ws, row, text):
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.fill, c.border, c.alignment = DET_LABEL, LABEL, BOX, DET_TOP


def det_value(ws, row, text, font=DET_VALUE, border=False):
    c = ws.cell(row=row, column=2, value=text)
    c.font, c.alignment = font, DET_TOP
    if border:
        c.border = BOX
    ws.row_dimensions[row].height = det_height(text)


def det_notes(d, x):
    notes = [f'- Defect ID: {d["id"]}', f'- Estado: {d["estado"]}', f'- Issue de GitHub: {d.get("issue") or "sin issue"}',
             f'- Caso de prueba: {d["caso"]}', f'- Encontrado en: {d["encontrado"]}', f'- Reportó: {d["reporter"]}',
             f'- Evidencia: {d["evidencia"]}']
    if x.get('correccion'):
        notes.append(f'- Corrección: {x["correccion"]}')
    if d.get('notas'):
        notes.append(f'- {d["notas"]}')
    return '\n'.join(notes)


def defect_sheet(wb, d, x):
    ws = wb.create_sheet(d['id'])
    ws.column_dimensions['A'].width = 31.83
    ws.column_dimensions['B'].width = 63.66
    row = 1
    det_label(ws, row, 'Title'); det_value(ws, row, f'[{d["id"]}] {d["titulo"]}', DET_TITLE, border=True); row += 2
    det_label(ws, row, 'Description / Overview'); det_value(ws, row, d['descripcion'], border=True); row += 2
    det_label(ws, row, 'Reproduction Steps')
    for i, paso in enumerate(d['pasos'], 1):
        det_value(ws, row, f'{i}- {paso}'); row += 1
    row += 1
    det_label(ws, row, 'Expected Behavior'); det_value(ws, row, d['esperado']); row += 2
    det_label(ws, row, 'Actual Behavior'); det_value(ws, row, d['obtenido']); row += 2
    det_label(ws, row, 'Incidence / Severity / Probability of reproduction')
    det_value(ws, row, f'Severidad {d["severidad"]} - {x["impacto"]}'); row += 1
    det_value(ws, row, x['reproduccion']); row += 2
    det_label(ws, row, 'Story and Acceptance Criteria affected'); det_value(ws, row, x['criterio']); row += 2
    det_label(ws, row, 'Browsers tested'); det_value(ws, row, x['navegadores']); row += 2
    det_label(ws, row, 'Environment'); det_value(ws, row, d['entorno']); row += 2
    det_label(ws, row, 'Notes'); det_value(ws, row, det_notes(d, x))


def gen_defects():
    detalle = json.loads((ROOT / 'fuentes' / 'defectos-detalle.json').read_text(encoding='utf-8'))
    wb = Workbook()
    ws = wb.active
    ws.title = 'Index'
    heads = ['Defect ID', 'Title', 'Severity', 'Status', 'Test case', 'User story', 'GitHub issue']
    ws.append(['Biyu – Entrega 1 – Defect reports']); ws['A1'].font = Font(bold=True, size=13)
    ws.append(heads)
    for i in range(1, len(heads) + 1):
        c = ws.cell(row=2, column=i); c.fill = HEAD; c.font = WHITE; c.border = BOX
    for d in DEFECTOS:
        ws.append([d['id'], d['titulo'], d['severidad'], d['estado'], d['caso'], d['historia'], d.get('issue') or 'sin issue'])
        for i in range(1, len(heads) + 1):
            x = ws.cell(row=ws.max_row, column=i); x.border = BOX; x.alignment = WRAP
        ws.cell(row=ws.max_row, column=1).hyperlink = f"#'{d['id']}'!A1"
    for col, w in zip('ABCDEFG', (11, 60, 10, 34, 26, 30, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'
    for d in DEFECTOS:
        defect_sheet(wb, d, detalle[d['id']])
    wb.save(ROOT / '04-reportes-de-defectos.xlsx')

    sev = Counter(d['severidad'] for d in DEFECTOS if not d['estado'].startswith('Cerrado'))
    L = ['# Proyecto Biyu – Entrega 1 · Reportes de defectos (V1)', '',
         'Planilla: `04-reportes-de-defectos.xlsx` (un índice y una hoja por defecto, con el formato "Defect Report - Template" de la cátedra). '
         'Cada defecto existe también como issue en GitHub con la etiqueta `bug`, salvo los nuevos de esta entrega, que quedan listos para cargar.', '',
         f'Escala de severidad y flujo de estados: `docs/07-plan-de-testing.md` §5. Severidad la fija quien reporta; **prioridad la fija el PO** '
         '(la que figura acá es la sugerida por quien reportó).', '',
         '## Resumen', '', '| ID | Título | Severidad | Prioridad sugerida | Estado | Caso | Historia |', '|---|---|---|---|---|---|---|']
    L += [f'| {d["id"]} | {d["titulo"]} | {d["severidad"]} | {d["prioridad"]} | {d["estado"]} | {d["caso"]} | {d["historia"]} |' for d in DEFECTOS]
    L += ['', '**Abiertos por severidad:** ' + ' · '.join(f'{k}: {sev.get(k, 0)}' for k in ('Crítica', 'Alta', 'Media', 'Baja')) +
          f' · Total abiertos: {sum(sev.values())}.', '']
    for d in DEFECTOS:
        L += [f'## {d["id"]} · {d["titulo"]}', '',
              '| Campo | Contenido |', '|---|---|',
              f'| Estado | {d["estado"]} |', f'| Severidad | {d["severidad"]} |', f'| Prioridad (sugerida) | {d["prioridad"]} |',
              f'| Encontrado en | {d["encontrado"]} |', f'| Caso de prueba | {d["caso"]} |', f'| Historia | {d["historia"]} |',
              f'| Issue | {d.get("issue") or "A crear"} |', f'| Reportó | {d["reporter"]} |', f'| Entorno | {d["entorno"]} |', '',
              d['descripcion'], '', '**Pasos para reproducir**', ''] + [f'{i}. {p}' for i, p in enumerate(d['pasos'], 1)] + [
              '', f'**Resultado esperado.** {d["esperado"]}', '', f'**Resultado obtenido.** {d["obtenido"]}', '',
              f'**Evidencia.** {d["evidencia"]}', '']
        if d.get('notas'):
            L += [f'**Notas.** {d["notas"]}', '']
    (ROOT / '04-reportes-de-defectos.md').write_text('\n'.join(L))


# ---------------------------------------------------------------- 05 reporte
def gen_report(extra):
    st, by_mod, by_prio, by_tipo, _ = metrics()
    total = len(CASOS)
    executed = st['PASSED'] + st['FAILED']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte de ejecución - Entrega1'
    ws.append(['Biyu - entrega 1 - Reporte de ejecución de test cases'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.merge_cells('A1:G1')
    heads = ['Test case ID', 'Prioridad', 'Título', 'User story', 'Camino feliz', 'Status', 'Defectos']
    ws.append(heads)
    for i in range(1, 8):
        c = ws.cell(row=2, column=i); c.fill = HEAD; c.font = WHITE; c.border = BOX
    for c in CASOS:
        ws.append([c['id'], c['prioridad'], c['titulo'], ', '.join(c['historias']), 'Sí' if c['feliz'] else 'No', status_of(c['id']), defects_of(c['id'])])
        for i in range(1, 8):
            x = ws.cell(row=ws.max_row, column=i); x.border = BOX; x.alignment = WRAP
        ws.cell(row=ws.max_row, column=6).fill = PatternFill('solid', fgColor=STATUS_FILL[status_of(c['id'])])
    for col, w in zip('ABCDEFG', (14, 10, 60, 18, 12, 12, 18)):
        ws.column_dimensions[col].width = w
    r0 = ws.max_row + 2
    ws.cell(row=r0, column=1, value='Resumen').font = BOLD
    summary = [('Planificados', total), ('Ejecutados', executed), ('PASSED', st['PASSED']), ('FAILED', st['FAILED']),
               ('BLOCKED', st['BLOCKED']), ('No ejecutados', st['NO EJECUTADO']), ('% avance (ejecutados/planificados)', pct(executed, total)),
               ('% aprobados sobre ejecutados', pct(st['PASSED'], executed))]
    for i, (k, v) in enumerate(summary, r0 + 1):
        ws.cell(row=i, column=1, value=k).fill = LABEL
        ws.cell(row=i, column=2, value=v)
    r1 = r0 + len(summary) + 2
    ws.cell(row=r1, column=1, value='Por módulo').font = BOLD
    ws.append([])
    for i, h in enumerate(['Módulo', 'Casos', 'PASSED', 'FAILED', 'BLOCKED'], 1):
        ws.cell(row=r1 + 1, column=i, value=h).fill = LABEL
    for j, (m, cnt) in enumerate(by_mod.items(), r1 + 2):
        for i, v in enumerate([f'{m} · {MODULOS[m]}', cnt['Total'], cnt['PASSED'], cnt['FAILED'], cnt['BLOCKED']], 1):
            ws.cell(row=j, column=i, value=v)
    wb.save(ROOT / '05-reporte-de-ejecucion.xlsx')
    return st, by_mod, by_prio, executed, total


if __name__ == '__main__':
    if sys.argv[1:] == ['defectos']:
        gen_defects()
        sys.exit()
    gen_spec()
    gen_exec()
    gen_defects()
    st, by_mod, by_prio, executed, total = gen_report(None)
    print(json.dumps({'total': total, 'executed': executed, **st, 'by_prio': {k: dict(v) for k, v in by_prio.items()},
                      'by_mod': {k: dict(v) for k, v in by_mod.items()}}, ensure_ascii=False))
